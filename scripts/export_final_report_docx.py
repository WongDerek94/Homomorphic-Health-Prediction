#!/usr/bin/env python3
"""Convert docs/FINAL_REPORT_DRAFT.md to docs/FINAL_REPORT.docx for submission."""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DRAFT = ROOT / "docs" / "FINAL_REPORT_DRAFT.md"
OUTPUT = ROOT / "docs" / "FINAL_REPORT.docx"
TEMPLATE = Path(
    "/home/utopia/BCIT_Courses/COMP 8047 - Major Project/"
    "Derek - Proposal Final Submission/Major Project Final Report Template.docx"
)
SCREENSHOTS_DIR = ROOT / "docs" / "screenshots"

SCREENSHOT_MAP = {
    "Gradio symptom selection UI": SCREENSHOTS_DIR / "01_symptom_selection.png",
    "Encryption step showing one-hot vector and ciphertext": SCREENSHOTS_DIR
    / "02_encryption_step.png",
    "Decrypted prediction output with top-3 diseases": SCREENSHOTS_DIR
    / "03_decrypted_top3.png",
}


def _ensure_screenshot_placeholders() -> None:
    """Create placeholder PNGs if real captures are not yet added."""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return

    SCREENSHOTS_DIR.mkdir(parents=True, exist_ok=True)
    labels = {
        "01_symptom_selection.png": "Symptom selection UI",
        "02_encryption_step.png": "Encryption step",
        "03_decrypted_top3.png": "Decrypted top-3 output",
    }
    for name, label in labels.items():
        path = SCREENSHOTS_DIR / name
        if path.is_file():
            continue
        img = Image.new("RGB", (1200, 675), color=(245, 247, 250))
        draw = ImageDraw.Draw(img)
        draw.rectangle([20, 20, 1180, 655], outline=(180, 190, 200), width=3)
        draw.text((60, 300), f"[INSERT SCREENSHOT: {label}]", fill=(90, 100, 110))
        img.save(path)


def _post_process_docx() -> None:
    doc = Document(str(OUTPUT))
    screenshot_re = re.compile(r"\[INSERT SCREENSHOT:\s*(.+?)\]", re.IGNORECASE)
    author_re = re.compile(r"\[INSERT AUTHOR (DISCLAIMER|USABILITY):", re.IGNORECASE)
    tbd_re = re.compile(r"\[TBD:", re.IGNORECASE)

    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue

        m = screenshot_re.search(text)
        if m:
            label = m.group(1).strip()
            path = next(
                (p for key, p in SCREENSHOT_MAP.items() if key.lower() in label.lower()),
                None,
            )
            paragraph.clear()
            if path and path.is_file():
                run = paragraph.add_run()
                run.add_picture(str(path), width=Inches(6.0))
            else:
                run = paragraph.add_run(f"[INSERT SCREENSHOT: {label}]")
                run.italic = True
                run.font.color.rgb = RGBColor(0x99, 0x66, 0x00)
            continue

        if author_re.search(text) or tbd_re.search(text):
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x00, 0x66, 0x99)

    doc.save(str(OUTPUT))


def main() -> int:
    if not DRAFT.is_file():
        print(f"Missing draft: {DRAFT}", file=sys.stderr)
        return 1

    _ensure_screenshot_placeholders()

    extra_args: list[str] = []
    if TEMPLATE.is_file():
        extra_args.extend(["--reference-doc", str(TEMPLATE)])

    pypandoc.convert_file(
        str(DRAFT),
        "docx",
        outputfile=str(OUTPUT),
        extra_args=extra_args,
    )
    _post_process_docx()
    print(f"Wrote {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
